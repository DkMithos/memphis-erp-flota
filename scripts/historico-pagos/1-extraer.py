# Extrae los pagos de "PAGOS ACUMULADO 2024 y 2025.xlsx" (hoja CAJA) y "Pagos acumulados 2026 con CC.xlsx" (hoja DATA)
# a un JSON normalizado para cargarlo en transacciones. Mapa de centros de costo Excel -> codigo ERP.
import openpyxl, warnings, json, datetime, re, collections
warnings.filterwarnings('ignore')
P = "C:/Users/URSULA/AppData/Local/Temp/claude/C--Users-URSULA-Proyectos-memphis-erp-flota/8d33990c-93a6-418b-ae32-d99f46f984d7/scratchpad/gg/pagos/"
F = "C:/Users/URSULA/AppData/Local/Temp/claude/C--Users-URSULA-Proyectos-memphis-erp-flota/8d33990c-93a6-418b-ae32-d99f46f984d7/scratchpad/gg/fin/"

CC = {
    'OFICINA CENTRAL': 'OFCENTRAL', 'OFICIAN CENTRAL': 'OFCENTRAL',
    'GORE.ICA': 'GOREICAPNP', 'GOREI': 'GOREICAPNP',
    'GOREC-2 PATRULLEROS': 'GCUSCOPNP', 'GORECZC': 'GCUSCOPNP',
    'CMSURCO': 'MSS-30', 'MSS-30': 'MSS-30',
    'GOREC-AMB': 'GCUSCOAMBU',
    'GORE.LORETO.BOMBEROS': 'GLOREBOMBE', 'GORE.IQUIT.BOMBEROS': 'GLOREBOMBE',
    'GORE-HUANUCO': 'GHUANUCOPNP', 'GORE_HUANUCO': 'GHUANUCOPNP',
    'MDI': 'MDI',
    'GORE-AMAZONAS': 'GAMAZONPNP', 'GORE_AMAZONAS_ PATRULLERO PNP': 'GAMAZONPNP',
    'MP CUSCO': 'MPCUSCOSERENAZGO',
    'TI': 'SISTEMAS',
    'SAN MIGUEL': 'MPSANMIGUEL',
    'LORETO AMB': 'LORETOAMB',
    'GOREC-HIDRO': 'GCUSCOHIDROAMB',
    'SELVA I': 'CMSELVAI',
    'DETRACCIONES': 'DETRACCIÓN',
    'INM PAN': 'INMPAN',
    'CAPITAL TRABAJO': 'PRÉSTAMOS',
    'COMERCIAL': 'MARKETING Y DESARROLLO', 'COMERCIAL HIDROAMBULANCIAS': 'MARKETING Y DESARROLLO',
    'C-OXI': 'C-OXI',
    'GORE-ANCASH': 'GANCASHAMB',
    'BRAVO 3': 'BRAVO 3',
    'FLORIDA': 'TERRENOS EEUU',
    'DEMO ISLA SEGURA': 'ISLASEGURA',
}
def cc(v):
    k = str(v or '').strip().upper()
    return CC.get(k), (str(v).strip() if v else None)

def fecha(v):
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    if isinstance(v, datetime.date): return v.isoformat()
    s = str(v or '').strip()
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})', s)
    if m: return f"{m[1]}-{m[2]}-{m[3]}"
    m = re.match(r'^(\d{2})[/-](\d{2})[/-](\d{4})', s)
    if m: return f"{m[3]}-{m[2]}-{m[1]}"
    return None

def num(v):
    if v is None or v == '': return None
    if isinstance(v, (int, float)): return round(float(v), 2)
    s = re.sub(r'[^0-9.,\-]', '', str(v))
    if not s: return None
    if ',' in s and s.rfind(',') > s.rfind('.'): s = s.replace('.', '').replace(',', '.')
    else: s = s.replace(',', '')
    try: return round(float(s), 2)
    except ValueError: return None

def txt(v, n=240):
    s = re.sub(r'\s+', ' ', str(v)).strip() if v is not None else ''
    return s[:n] or None

out = []; sin_cc = collections.Counter()
# ── 2026 ─────────────────────────────────────────────────────────────────────
ws = openpyxl.load_workbook(P + "Pagos acumulados 2026 con CC.xlsx", data_only=True, read_only=True)['DATA']
rows = list(ws.iter_rows(values_only=True))
for r in rows[2:]:
    if not r or r[0] is None: continue
    item, ccx, cat, tipo, comp, razon, desc, pen, usd, ope, fpago, mes, tipo_egreso = (list(r) + [None]*13)[:13]
    f = fecha(fpago)
    if not f: continue
    codigo, etiqueta = cc(ccx)
    if etiqueta and not codigo: sin_cc[etiqueta] += 1
    cambio = str(tipo_egreso or '').strip().upper() == 'CAMBIO DE MONEDA' or str(cat or '').strip().upper() == 'CAMBIO DE MONEDA'
    for moneda, monto in (('PEN', num(pen)), ('USD', num(usd))):
        if not monto: continue
        out.append({
            'origen': 'acumulado_2026', 'fila': int(item), 'fecha': f, 'moneda': moneda, 'monto': abs(monto),
            'tipo': ('ingreso' if (cambio and moneda == 'USD') or monto < 0 else 'egreso'),
            'cuenta': '806' if moneda == 'PEN' else '830',
            'cc': codigo, 'cc_excel': etiqueta, 'categoria': txt(cat, 60) or 'Sin categoría',
            'tipo_doc': txt(tipo, 40), 'comprobante': txt(comp, 40), 'proveedor': txt(razon, 160), 'descripcion': txt(desc),
            'ope': txt(ope, 20), 'tc': None, 'detraccion': None, 'retencion': None, 'cambio_moneda': cambio,
        })
# ── 2024-2025 ─────────────────────────────────────────────────────────────────
ws = openpyxl.load_workbook(F + "14_pagos_acumulado_2024_2025.xlsx", data_only=True, read_only=True)['CAJA']
rows = list(ws.iter_rows(values_only=True))
for r in rows[1:]:
    if not r or r[0] is None: continue
    (l, ccx, cat, fcomp, tipo, comp, razon, desc, pen, usd, det, ret, tc, transf, ope, fpago, fdet) = (list(r) + [None]*17)[:17]
    f = fecha(fpago)
    if not f: continue
    codigo, etiqueta = cc(ccx)
    if etiqueta and not codigo: sin_cc[etiqueta] += 1
    for moneda, monto in (('PEN', num(pen)), ('USD', num(usd))):
        if not monto: continue
        out.append({
            'origen': 'acumulado_2024_2025', 'fila': int(l) if isinstance(l, (int, float)) else None, 'fecha': f, 'moneda': moneda, 'monto': abs(monto),
            'tipo': 'ingreso' if monto < 0 else 'egreso',
            'cuenta': '545' if moneda == 'PEN' else '553',
            'cc': codigo, 'cc_excel': etiqueta, 'categoria': txt(cat, 60) or 'Sin categoría',
            'tipo_doc': txt(tipo, 40), 'comprobante': txt(comp, 40), 'proveedor': txt(razon, 160), 'descripcion': txt(desc),
            'ope': txt(ope, 20), 'tc': num(tc), 'detraccion': num(det), 'retencion': num(ret), 'cambio_moneda': False,
        })
json.dump(out, open(F + "historico_pagos.json", "w", encoding='utf-8'), ensure_ascii=False)
tot = collections.Counter(); montos = collections.defaultdict(float)
for o in out:
    tot[(o['origen'], o['moneda'], o['tipo'])] += 1; montos[(o['origen'], o['moneda'], o['tipo'])] += o['monto']
print("filas:", len(out))
for k in sorted(tot): print("  ", k, tot[k], round(montos[k]))
print("CC sin mapa:", dict(sin_cc))
print("sin CC (vacío):", sum(1 for o in out if not o['cc_excel']))
print("con OPE:", sum(1 for o in out if o['ope']))
